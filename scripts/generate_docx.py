from docx import Document
from docx.shared import Pt

doc = Document()

doc.add_heading('Darukaa Hackathon Submission', level=1)

doc.add_heading('Submission Header', level=2)
doc.add_paragraph('Name: Nikhat Mulla')
doc.add_paragraph('Position: Full Stack Developer Internship (Climate-Tech & Nature Intelligence)')
doc.add_paragraph('Date: September 2026')

doc.add_heading('Links', level=2)
doc.add_paragraph('GitHub Repository URL: https://github.com/<your-username>/darukaa-earth')
doc.add_paragraph('Live Application URL: https://darukaa-earth.vercel.app')
doc.add_paragraph('Demo Credentials: email: test@darukaa.earth, password: Test1234')

doc.add_heading('Brief Architecture Overview', level=2)
doc.add_paragraph('- Frontend: React + Mapbox GL JS + Chart.js')
doc.add_paragraph('- Backend: FastAPI (Python) with JWT Auth')
doc.add_paragraph('- Database: PostgreSQL with PostGIS on Supabase')
doc.add_paragraph('- CI/CD: Pre-commit formatting via Husky and automated build validation via GitHub Actions')

doc.add_heading('Local Setup Instructions', level=2)
doc.add_paragraph('Backend:')
doc.add_paragraph('1. Create and activate virtualenv:')
doc.add_paragraph('   python -m venv venv')
doc.add_paragraph('   venv\\Scripts\\activate')
doc.add_paragraph('2. Install dependencies: pip install -r backend/requirements.txt')
doc.add_paragraph('3. Run: cd backend && uvicorn main:app --reload')

doc.add_paragraph('\nFrontend:')
doc.add_paragraph('1. cd frontend && npm install')
doc.add_paragraph('2. npm run dev')

doc.add_heading('Collaborator Access', level=2)
doc.add_paragraph('If the repo is private, please add these collaborators under Settings > Collaborators:')
doc.add_paragraph('- ankita.dasgupta@darukaa.com')
doc.add_paragraph('- harsh.kumar@darukaa.com')
doc.add_paragraph('- utkarsh.gauniyal@darukaa.com')
doc.add_paragraph('- guneet.mutreja@darukaa.com')

output_path = '../Darukaa_Hackathon_Submission.docx'
print('Saving doc to', output_path)
doc.save(output_path)
print('Done')
